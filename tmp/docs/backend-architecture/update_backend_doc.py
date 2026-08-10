from copy import deepcopy
from pathlib import Path
import shutil

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor


ROOT = Path(__file__).resolve().parents[3]
REFERENCE = ROOT / "backend" / "docs" / "LearnerOS_Backend_Architecture.docx"
DRAFT = ROOT / "tmp" / "docs" / "backend-architecture" / "LearnerOS_Backend_Architecture.updated.docx"


def copy_rpr(run):
    return deepcopy(run._r.rPr) if run is not None and run._r.rPr is not None else None


def apply_rpr(run, rpr):
    if rpr is not None:
        if run._r.rPr is not None:
            run._r.remove(run._r.rPr)
        run._r.insert(0, deepcopy(rpr))


def find_para(doc, text):
    for para in doc.paragraphs:
        if para.text == text:
            return para
    raise KeyError(f"paragraph not found: {text[:90]}")


def set_para(para, text):
    rpr = copy_rpr(para.runs[0] if para.runs else None)
    para.clear()
    run = para.add_run(text)
    apply_rpr(run, rpr)


def set_labeled_para(para, label, body):
    label_rpr = copy_rpr(para.runs[0] if para.runs else None)
    body_rpr = copy_rpr(para.runs[-1] if para.runs else None)
    para.clear()
    r1 = para.add_run(label)
    apply_rpr(r1, label_rpr)
    r1.bold = True
    r2 = para.add_run(body)
    apply_rpr(r2, body_rpr)
    r2.bold = False


def set_cell(cell, text):
    para = cell.paragraphs[0]
    rpr = copy_rpr(para.runs[0] if para.runs else None)
    para.clear()
    run = para.add_run(text)
    apply_rpr(run, rpr)
    for extra in list(cell.paragraphs[1:]):
        extra._element.getparent().remove(extra._element)


def set_labeled_cell(cell, label, body):
    para = cell.paragraphs[0]
    label_rpr = copy_rpr(para.runs[0] if para.runs else None)
    body_rpr = copy_rpr(para.runs[-1] if para.runs else None)
    para.clear()
    r1 = para.add_run(label)
    apply_rpr(r1, label_rpr)
    r1.bold = True
    r2 = para.add_run(body)
    apply_rpr(r2, body_rpr)
    r2.bold = False
    for extra in list(cell.paragraphs[1:]):
        extra._element.getparent().remove(extra._element)


shutil.copy2(REFERENCE, DRAFT)
doc = Document(DRAFT)

# Cover and guide
set_para(
    find_para(doc, "Routing, APIs, persistence, LLM orchestration, caching, telemetry, evaluation, and deployment"),
    "Routing, APIs, persistence, assessment context, insight lifecycle, LLM orchestration, caching, and deployment",
)
set_para(
    find_para(doc, "12. Telemetry, production assessments, and evaluations"),
    "12. Assessment context, concept fallback, and insight scope",
)
set_cell(
    doc.tables[0].cell(3, 1),
    "Code-grounded implementation reference; assessment behavior updated after the subsection/concept-scope changes",
)
set_labeled_cell(
    doc.tables[1].cell(0, 0),
    "Reading guide\n",
    "This document describes the backend currently represented by the source tree. Assessment behavior is grounded in the production router and its focused regression tests. Unrelated evaluation, export, promotion, and observability scripts are outside the scope of this revision.",
)

# System role
set_labeled_para(
    find_para(doc, "Generation layer: question/MCQ generation, assessment evaluation, explanation generation, insight reconciliation, embeddings, and tutor responses."),
    "Generation layer: ",
    "written-question and MCQ generation are target-subsection scoped; evaluation uses the target subsection plus full-section reference context and then creates validated, concept-linked insights.",
)

# Runtime lifecycle table
set_cell(doc.tables[3].cell(3, 1), "Neo4j read of section/subsection content plus direct-section or chapter-fallback concept candidates")
set_cell(doc.tables[3].cell(3, 2), "Missing section content returns 404; a subsection outside the section returns 400; zero valid concept candidates is allowed")

# Routing description and assessment row
set_para(
    find_para(doc, "The route registration source of truth is backend/app/main.py. Router modules group endpoints by domain. Path parameters that contain Neo4j identifiers use the :path converter because IDs contain colons and other separators."),
    "The route registration source of truth is backend/app/main.py. Router modules group endpoints by domain. Most Neo4j identifiers use the :path converter. The section-concepts route deliberately uses /sections/{section_id}/concepts without :path so it cannot absorb the /sections/{id}/test/... namespace; section IDs contain colons but no slash.",
)
set_cell(doc.tables[4].cell(6, 2), "Remediation and insight-specific testing")
set_cell(doc.tables[4].cell(7, 2), "Target-subsection question/MCQ generation; personalized evaluation with validated insight scope")

# Curriculum model and insight flow
set_para(
    find_para(doc, "Section and/or subsection nodes connect to Concept nodes through REQUIRES and teaches/related relationships used by graph and prerequisite views."),
    "Sections connect to Concept nodes through REQUIRES. Assessment first uses concepts directly required by the active section; if none exist, it collects concepts required by sections in the same chapter as evaluation-only fallback candidates.",
)
set_para(
    find_para(doc, "Insights are student-specific evidence attached to curriculum concepts and sources. The active state is a query projection over historical evidence, not a replacement for historical records. The normal MCQ flow is: answer -> model evaluation -> structured evidence -> insight creation -> reconciliation -> embedding -> persistence."),
    "Insights are student-specific evidence attached to curriculum concepts and sources. The active state is a query projection over historical evidence, not a replacement for historical records. The written-answer and MCQ flows are: resolve target subsection -> evaluate against subsection content -> select valid concept candidates -> validate structured evidence -> insight creation -> reconciliation -> embedding -> persistence.",
)
set_cell(
    doc.tables[8].cell(0, 0),
    "Student answer\n"
    "  -> resolve and validate the requested subsection inside the section\n"
    "  -> evaluate against target subsection content; full section is reference context only\n"
    "  -> use direct section concepts, or chapter concepts when the section has none\n"
    "  -> reject insight IDs outside the resolved candidate set\n"
    "  -> attach accepted evidence to the target subsection/source and concept\n"
    "  -> reconcile against active insight(s) for the same student + source/concept key\n"
    "  -> preserve history, embed the active insight, and retrieve it later for teaching/tutor context",
)

# Prompt paths table
set_cell(doc.tables[12].cell(1, 1), "Target subsection content as primary scope; complete parent section as definitions/surrounding reference only; no concept candidates in the generation prompt")
set_cell(doc.tables[12].cell(1, 2), "Non-personalized and share-cache eligible; requested subsection ID is authoritative")
set_cell(doc.tables[12].cell(2, 1), "Question/answer, target subsection ground truth, full-section reference context, and direct-section or chapter-fallback concept candidates")
set_cell(doc.tables[12].cell(2, 2), "Personalized evaluation; accepted insights must use an exact candidate concept ID and the target subsection source ID")

# Repair spacing in the retained production-operation bullets.
for old, label, body in [
    ("generate_textaccepts system/developer/user messages, optional image data URLs, optional JSON mode, model/provider selection, and operation labels.", "generate_text: ", "accepts system/developer/user messages, optional image data URLs, optional JSON mode, model/provider selection, and operation labels."),
    ("generate_jsoncalls the model, parses JSON, validates the expected shape in the caller, and retries parse/generation failures.", "generate_json: ", "calls the model, parses JSON, validates the expected shape in the caller, and retries parse/generation failures."),
    ("embedcreates an embedding for insight or query text through Fireworks; the embedding operation is separated from generation purpose.", "embed: ", "creates an embedding for insight or query text through Fireworks; the embedding operation is separated from generation purpose."),
    ("Provider fallbackis not universal. It is used only when the caller supplies fallback targets; a provider misconfiguration can therefore fail the operation rather than silently switching models.", "Provider fallback: ", "is not universal. It is used only when the caller supplies fallback targets; a provider misconfiguration can therefore fail the operation rather than silently switching models."),
    ("Content captureprompt/response/image content is fingerprinted or hidden in telemetry by default; full content capture requires OTEL_CAPTURE_CONTENT=true.", "Content capture: ", "prompt/response/image content is fingerprinted or hidden in telemetry by default; full content capture requires OTEL_CAPTURE_CONTENT=true."),
]:
    set_labeled_para(find_para(doc, old), label, body)

# Generation cache identity no longer contains a concept parameter.
set_para(
    find_para(doc, "The deterministic key includes endpoint type, section ID, subsection ID, concept ID where relevant, variant, generation model, and prompt schema version. It deliberately excludes auth identity. Bumping GEN_PROMPT_VERSION or changing the model naturally creates a new key namespace."),
    "The deterministic question/MCQ generation key includes endpoint type/version, section ID, requested subsection ID, variant, provider/model/fallback targets, and prompt schema version. It deliberately excludes auth identity and concept ID because concepts do not drive question generation. The v2 endpoint namespace prevents older, broader prompt results from being reused.",
)

# Replace Section 12 in place with the updated production assessment contract.
set_para(
    find_para(doc, "12. Telemetry, Production Assessments, And Evaluations"),
    "12. Assessment Context, Concept Fallback, And Insight Scope",
)
set_para(
    find_para(doc, "The current observability path is OpenTelemetry plus structured logs, not MLflow. OpenTelemetry can instrument FastAPI, OpenAI-compatible calls, spans, and metrics. It exports through OTLP HTTP when enabled. The backend also has assessment-specific stage timing and lifecycle events."),
    "The production assessment contract separates question content from insight classification. Question and MCQ generation are controlled by the requested subsection content, with the full parent section supplied only as reference context. Concept candidates are introduced only during answer evaluation and insight creation.",
)

assessment_rows = [
    ("STAGE", "CONTENT BOUNDARY", "CONCEPT BEHAVIOR"),
    ("Metadata load", "Collect ordered subsection IDs, titles, and content; build complete section reference text", "Select concepts directly required by the section; if none, collect concepts required by sections in the same chapter"),
    ("Target resolution", "Require the submitted subsection ID to belong to the requested section", "No concept is needed to resolve or render the subsection; invalid membership returns 400"),
    ("Question / MCQ generation", "Target subsection content is primary; full section may clarify notation, definitions, or surrounding context", "Concept IDs and names are excluded from the generation prompt and cache identity"),
    ("Answer evaluation", "Judge only against the target subsection; do not require sibling-only facts", "Expose either direct-section candidates or explicitly labeled chapter-fallback candidates"),
    ("Insight validation", "Accept only evidence meaningfully supported by the tested question and learner answer", "Reject concept IDs outside the resolved candidate set and invalid type/category values; an empty insight list is valid"),
    ("Persistence", "Force source_id to the validated target subsection", "Persist only accepted concept links; unauthenticated or empty-insight cases are skipped without changing the evaluation result"),
]
for r, row in enumerate(assessment_rows):
    for c, value in enumerate(row):
        set_cell(doc.tables[22].cell(r, c), value)

# Keep the replacement table visually consistent and prevent rows splitting.
for cell in doc.tables[22].rows[0].cells:
    for run in cell.paragraphs[0].runs:
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
for row in doc.tables[22].rows:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)

set_para(find_para(doc, "Where traces and assessments go"), "Concept-scope selection")
scope_bullets = [
    "Direct scope: if the active section has one or more direct REQUIRES Concept relationships, only those concepts are candidates for assessment insights.",
    "Chapter fallback: if the active section has no direct concepts, concepts required by other sections in the same chapter become evaluation-only candidates. This is the section 10.3 case.",
    "No candidates: if neither direct nor chapter concepts exist, the answer can still be evaluated, but returned insights are filtered to an empty list.",
    "Source integrity: every accepted insight receives the validated target subsection as source_id, regardless of what the model returns.",
]
for old, new in zip([
    "OTel traces and metrics go to the configured OTLP exporter endpoint, such as an observability backend or collector. The backend does not persist all telemetry in Neo4j.",
    "Structured request and assessment events go to process/container logs and can be collected by Fly or a log sink.",
    "AssessmentAttempt records and insight state go to Neo4j. These are product records, not an observability substitute.",
    "Production eval code in backend/evals uses the same LLMService and contracts, while endpoint-focused tests exercise the production assessment path in backend/tests.",
], scope_bullets):
    set_para(find_para(doc, old), new)

set_para(find_para(doc, "Telemetry configuration"), "Focused regression coverage")
set_para(
    find_para(doc, "Environment configuration shape"),
    "Current contract checks in backend/tests/test_assessment_context.py",
)
set_cell(
    doc.tables[23].cell(0, 0),
    "generation uses target content plus full-section reference and excludes concepts\n"
    "subsection IDs outside the section are rejected before model evaluation\n"
    "a section with no valid concepts evaluates successfully with empty insights\n"
    "MCQ evaluation accepts an insight from the declared chapter-fallback scope\n"
    "the curriculum concept route does not shadow the /test namespace\n"
    "verified locally: 7 focused tests passed",
)
set_labeled_cell(
    doc.tables[24].cell(0, 0),
    "Important boundary\n",
    "The full section is sent because an individual subsection may omit definitions or nearby context. It is reference material, not permission to test sibling subsections. Likewise, chapter concepts are fallback labels for evaluation and insight creation; they do not broaden question generation.",
)

# Current caveat and hardening language
set_cell(doc.tables[28].cell(4, 0), "Chapter concept fallback")
set_cell(doc.tables[28].cell(4, 1), "Fallback candidates can be broader than direct section links")
set_cell(doc.tables[28].cell(4, 2), "Strict prompt and exact-ID validation are required; empty insights are a valid outcome")
set_para(
    find_para(doc, "Add a schema/relationship audit job for concept-to-section, section-to-chapter, grade-to-subject, and cluster membership consistency."),
    "Add a schema/relationship audit for direct section concepts and chapter-fallback coverage, including sections that intentionally have no concept links.",
)
set_para(
    find_para(doc, "Add end-to-end production-like tests that start FastAPI with test doubles only at external boundaries while exercising the actual router and service path."),
    "Keep production-like route tests for both direct-concept and chapter-fallback sections, including the no-candidate case and cache-version changes.",
)
set_labeled_cell(
    doc.tables[29].cell(0, 0),
    "Source files\n",
    "Primary implementation sources referenced by this document include backend/app/main.py, auth.py, config.py, database.py, routers/curriculum.py, routers/test.py, routers/tutor.py, services/llm.py, services/generation_cache.py, services/semantic_clustering.py, services/clustering_runner.py, teacher_analytics.py, and backend/tests/test_assessment_context.py. Unrelated eval/export/promotion/observability scripts are not used as behavioral evidence in this revision.",
)

# Word uses this metadata for assistive navigation and repeated table headings.
# Single-row callout/layout tables receive the marker as their first row too.
for table in doc.tables:
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        tbl_header = OxmlElement("w:tblHeader")
        tbl_header.set(qn("w:val"), "true")
        tr_pr.append(tbl_header)

doc.save(DRAFT)
print(DRAFT)

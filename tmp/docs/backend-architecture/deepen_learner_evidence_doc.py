from copy import deepcopy
from pathlib import Path
import shutil

from docx import Document


ROOT = Path(__file__).resolve().parents[3]
REFERENCE = ROOT / "backend" / "docs" / "LearnerOS_Backend_Architecture.docx"
DRAFT = ROOT / "tmp" / "docs" / "backend-architecture" / "LearnerOS_Backend_Architecture.deepened.docx"


def replace_paragraph(doc: Document, old: str, new: str) -> None:
    for paragraph in doc.paragraphs:
        if paragraph.text == old:
            properties = deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs and paragraph.runs[0]._r.rPr is not None else None
            paragraph.clear()
            run = paragraph.add_run(new)
            if properties is not None:
                run._r.insert(0, properties)
            return
    raise KeyError(old)


def replace_cell(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    properties = deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs and paragraph.runs[0]._r.rPr is not None else None
    paragraph.clear()
    run = paragraph.add_run(text)
    if properties is not None:
        run._r.insert(0, properties)
    for extra in list(cell.paragraphs[1:]):
        extra._element.getparent().remove(extra._element)


shutil.copy2(REFERENCE, DRAFT)
doc = Document(DRAFT)

replace_paragraph(
    doc,
    "Assessment Execution transactionally creates durable Neo4j LearnerEvidenceJob records before reporting evidence as queued. The deep persistence module recovers interrupted jobs and owns reconciliation, embedding, idempotent insight writes, retry/dead-letter behavior, cache invalidation, and atomic Assessment Attempt result accounting; routers only call its stable interface.",
    "Assessment Execution transactionally creates durable Neo4j LearnerEvidenceJob records before reporting evidence as queued. Indexed lock-before-predicate claims issue one worker a token and lease; the deep module recovers expired or unaccounted jobs and owns reconciliation, embedding, idempotent insight writes, dead letters, cache invalidation, and atomic Assessment Attempt accounting.",
)
replace_paragraph(
    doc,
    "Learner Evidence is student-specific and attached to curriculum concepts and sources. Current Learner State is a projection over historical evidence. Assessment Execution validates evidence, then durable LearnerEvidenceJob processing owns reconciliation, embedding, idempotent graph persistence, supersession, recovery, cache invalidation, dead-letter recording, and attempt completion accounting.",
    "Learner Evidence is student-specific and attached to curriculum concepts and sources. Current Learner State is a projection over historical evidence. Assessment Execution validates evidence, then leased LearnerEvidenceJob processing owns reconciliation, embedding, idempotent graph persistence, supersession, recovery, cache invalidation, dead-letter recording, and attempt completion accounting.",
)

replace_cell(
    doc.tables[3].cell(5, 1),
    "Assessment Attempt queueing plus durable LearnerEvidenceJob creation in one Neo4j write; indexed token/lease claims coordinate workers, while the module owns recovery, idempotent insight writes, dead letters, cache invalidation, and atomic result counters",
)
replace_cell(
    doc.tables[3].cell(5, 2),
    "The learner-facing evaluation remains successful when asynchronous persistence fails; queued and unaccounted jobs recover after restarts, while terminal failures remain observable",
)
replace_cell(
    doc.tables[8].cell(0, 0),
    "Student answer\n"
    "  -> Assessment Execution resolves context, evaluates, and validates Learner Evidence\n"
    "  -> module transactionally records durable LearnerEvidenceJob work before returning queued\n"
    "  -> one worker acquires an indexed claim token and lease through the stable interface\n"
    "  -> module reconciles the same student + source + concept key\n"
    "  -> module embeds and persists the new active insight with supersession history\n"
    "  -> module invalidates derived views and atomically completes the job plus attempt counters\n"
    "  -> UI polls the learner-owned attempt until terminal and reloads Current Learner State",
)
replace_cell(
    doc.tables[22].cell(6, 1),
    "Force source_id to the validated target subsection, then submit Learner Evidence through the stable persistence interface",
)
replace_cell(
    doc.tables[22].cell(6, 2),
    "One deep module owns reconciliation, embedding, Neo4j persistence, retries, dead letters, cache invalidation, attempt accounting, and terminal convergence",
)
replace_cell(
    doc.tables[25].cell(4, 1),
    "Assessment Execution and Learner Evidence persistence emit lifecycle events through the same assessment ID; persistence stages remain internal to the deep module",
)
replace_cell(
    doc.tables[25].cell(4, 2),
    "One assessment can be followed across synchronous evaluation and the complete background Current Learner State convergence path",
)

replace_cell(
    doc.tables[31].cell(0, 0),
    "Source files\nPrimary sources include backend/app/{main,auth,config,database,observability,telemetry,assessment_observability}.py; routers/{curriculum,test,insights,tutor}.py; services/{llm,generation_cache,assessment_execution,learner_evidence,assessment_attempts,semantic_clustering,clustering_runner}.py; teacher_analytics.py; and focused tests under backend/tests/.",
)

doc.save(DRAFT)
print(DRAFT)

from copy import deepcopy
from pathlib import Path
import shutil

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor
from docx.table import Table


ROOT = Path(__file__).resolve().parents[3]
SOURCE = ROOT / "backend" / "docs" / "LearnerOS_Backend_Architecture.docx"
DRAFT = ROOT / "tmp" / "docs" / "backend-architecture" / "LearnerOS_Backend_Architecture.observability.docx"


def find_para(doc, text):
    for para in doc.paragraphs:
        if para.text == text:
            return para
    raise KeyError(text)


def copy_rpr(run):
    return deepcopy(run._r.rPr) if run is not None and run._r.rPr is not None else None


def apply_rpr(run, rpr):
    if rpr is not None:
        if run._r.rPr is not None:
            run._r.remove(run._r.rPr)
        run._r.insert(0, deepcopy(rpr))


def set_para(para, text):
    rpr = copy_rpr(para.runs[0] if para.runs else None)
    para.clear()
    run = para.add_run(text)
    apply_rpr(run, rpr)


def set_cell(cell, text):
    para = cell.paragraphs[0]
    rpr = copy_rpr(para.runs[0] if para.runs else None)
    para.clear()
    run = para.add_run(text)
    apply_rpr(run, rpr)
    for extra in list(cell.paragraphs[1:]):
        extra._element.getparent().remove(extra._element)


def set_labeled_cell(cell, label, body):
    para = cell.paragraphs[0]
    label_rpr = copy_rpr(para.runs[0] if para.runs else None)
    body_rpr = copy_rpr(para.runs[-1] if para.runs else None)
    para.clear()
    lead = para.add_run(label)
    apply_rpr(lead, label_rpr)
    lead.bold = True
    text = para.add_run(body)
    apply_rpr(text, body_rpr)
    text.bold = False
    for extra in list(cell.paragraphs[1:]):
        extra._element.getparent().remove(extra._element)


def insert_paragraph_before(doc, anchor, text, style):
    para = doc.add_paragraph(text, style=style)
    anchor._p.addprevious(para._p)
    return para


def insert_table_before(doc, anchor, template_table, rows):
    tbl_xml = deepcopy(template_table._tbl)
    anchor._p.addprevious(tbl_xml)
    table = Table(tbl_xml, doc._body)
    while len(table.rows) < len(rows):
        table._tbl.append(deepcopy(table.rows[-1]._tr))
    while len(table.rows) > len(rows):
        table._tbl.remove(table.rows[-1]._tr)
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            set_cell(table.cell(r, c), value)
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
    return table


def insert_callout_before(doc, anchor, template_table, label, body):
    tbl_xml = deepcopy(template_table._tbl)
    anchor._p.addprevious(tbl_xml)
    table = Table(tbl_xml, doc._body)
    cell = table.cell(0, 0)
    para = cell.paragraphs[0]
    label_rpr = copy_rpr(para.runs[0] if para.runs else None)
    body_rpr = copy_rpr(para.runs[-1] if para.runs else None)
    para.clear()
    lead = para.add_run(label)
    apply_rpr(lead, label_rpr)
    lead.bold = True
    text = para.add_run(body)
    apply_rpr(text, body_rpr)
    text.bold = False
    return table


shutil.copy2(SOURCE, DRAFT)
doc = Document(DRAFT)
config_table = doc.tables[26]
source_table = doc.tables[29]
assessment_callout_table = doc.tables[24]

# Cover, reading guide, and contents.
set_para(
    doc.paragraphs[2],
    "Routing, APIs, persistence, assessment context, insight lifecycle, LLM orchestration, caching, observability, and deployment",
)
set_labeled_cell(
    doc.tables[1].cell(0, 0),
    "Reading guide\n",
    "This document describes the backend currently represented by the source tree. Assessment behavior is grounded in the production router and its focused regression tests. Observability behavior is grounded in the request middleware, OpenTelemetry initializer, LLM service, assessment lifecycle helper, durable assessment records, and their tests. Unrelated evaluation, export, and promotion scripts are outside the behavioral scope of this revision.",
)
contents_security = find_para(doc, "13. Security and privacy controls")
insert_paragraph_before(doc, contents_security, "13. Observability and OpenTelemetry", "List Bullet")
set_para(contents_security, "14. Security and privacy controls")
set_para(find_para(doc, "14. Configuration and deployment"), "15. Configuration and deployment")
set_para(find_para(doc, "15. Operational runbook and known gaps"), "16. Operational runbook and known gaps")

# Renumber the existing final sections and insert a dedicated observability section.
security_heading = find_para(doc, "13. Security And Privacy Controls")
set_para(security_heading, "14. Security And Privacy Controls")
set_para(find_para(doc, "14. Configuration And Deployment"), "15. Configuration And Deployment")
set_para(find_para(doc, "15. Operational Runbook And Known Gaps"), "16. Operational Runbook And Known Gaps")

insert_paragraph_before(doc, security_heading, "13. Observability And OpenTelemetry", "Heading 1")
insert_paragraph_before(
    doc,
    security_heading,
    "LearnerOS combines always-available structured application logs with optional OpenTelemetry tracing and metrics. Request IDs correlate HTTP activity, LLM calls, assessment lifecycle stages, and background persistence. When enabled, telemetry is exported over OTLP HTTP to Phoenix or another compatible collector; the product database remains the source of truth for assessment records and insights.",
    "Normal",
)
insert_paragraph_before(doc, security_heading, "Signal and correlation model", "Heading 2")
for text in [
    "Request correlation: middleware accepts a safe X-Request-ID matching [A-Za-z0-9_-]{8,128}, otherwise generates an opaque UUID-derived value. The ID is stored in a context variable, returned in the response header, and added to structured events and custom spans.",
    "Structured logs: trace_event emits compact JSON with UTC timestamp, event, request_id, and bounded fields. trace_exception adds error type and a message capped at 500 characters while retaining the traceback in Fly logs.",
    "Assessment correlation: each answer receives an opaque assessment:<uuid> identifier. Lifecycle events use this ID across context loading, model work, validation, queueing, background persistence, retries, dead-letter recording, and completion.",
]:
    insert_paragraph_before(doc, security_heading, text, "List Bullet")

insert_paragraph_before(doc, security_heading, "Export pipeline and instrumentation", "Heading 2")
observability_rows = [
    ("SIGNAL / PATH", "IMPLEMENTATION", "OPERATING EFFECT"),
    ("HTTP requests", "FastAPI instrumentation plus request-start, request-complete, latency, status, and unhandled-exception JSON events; /health is excluded from auto-instrumentation", "Route spans and searchable Fly logs share request correlation without health-check noise"),
    ("LLM generation", "Custom llm.<operation> spans record provider, model, operation, input/output sizes, fingerprints, image count, latency, token usage when available, fallbacks, retries, and exceptions", "Separates model latency and provider failure from the surrounding route"),
    ("OpenAI-compatible calls", "OpenInference OpenAI instrumentation is installed on the shared client path", "Provides provider-call spans compatible with Phoenix and other OTLP backends"),
    ("Assessment lifecycle", "assessment.lifecycle log events and span events use scalar attributes; the active span receives assessment.id and assessment.status", "One assessment can be followed through synchronous evaluation and background insight persistence"),
    ("Assessment metrics", "learneros.assessment.stage.duration histogram records milliseconds by bounded stage, assessment kind, result, and optional model", "Supports latency and failure analysis without student, concept, or assessment IDs as metric labels"),
    ("OTLP export", "BatchSpanProcessor with OTLPSpanExporter; PeriodicExportingMetricReader with OTLPMetricExporter; shared URL-decoded headers; parent-based ratio sampling", "Traces batch asynchronously; metrics export at the configured interval to an explicit metrics endpoint or a /v1/metrics endpoint derived from /v1/traces"),
    ("Shutdown", "Trace and metric providers force-flush with a 5-second timeout and then shut down during application lifespan cleanup", "Reduces loss of buffered telemetry during a controlled stop"),
]
observability_table = insert_table_before(doc, security_heading, config_table, observability_rows)
for cell in observability_table.rows[0].cells:
    for run in cell.paragraphs[0].runs:
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

insert_paragraph_before(doc, security_heading, "Configuration and privacy defaults", "Heading 2")
for text in [
    "Enablement: OTEL_ENABLED defaults to false, making configure_telemetry a no-op. OTEL_SERVICE_NAME and ENVIRONMENT become resource attributes when enabled.",
    "Destinations and cadence: OTEL_EXPORTER_OTLP_ENDPOINT configures traces; OTEL_EXPORTER_OTLP_METRICS_ENDPOINT can override metrics; OTEL_METRIC_EXPORT_INTERVAL_MS defaults to 60,000 ms; OTEL_SAMPLE_RATIO is constrained to 0.0-1.0.",
    "Authentication headers: OTEL_EXPORTER_OTLP_HEADERS accepts comma-separated key=value pairs and URL-decodes values, allowing values such as Authorization=Bearer%20<token> without placing credentials in source code.",
    "Content protection: OTEL_CAPTURE_CONTENT defaults to false. The initializer sets OpenInference flags to hide inputs, outputs, and input images; LearnerOS custom spans/logs use lengths and SHA-256-derived 16-character fingerprints instead of prompt or response bodies.",
]:
    insert_paragraph_before(doc, security_heading, text, "List Bullet")
insert_callout_before(
    doc,
    security_heading,
    assessment_callout_table,
    "Privacy boundary\n",
    "Structured assessment logs may still include operational identifiers such as section, concept, source, insight, and assessment IDs. Metric labels intentionally exclude high-cardinality student, concept, and assessment identifiers. Durable AssessmentAttempt records can contain question and text-answer content for authenticated students; image answers are represented by SHA-256 hashes rather than image data.",
)

insert_paragraph_before(doc, security_heading, "Durable assessment operations and verification", "Heading 2")
for text in [
    "AssessmentAttempt nodes persist evaluation status, model/fallback state, scores or correctness, insight queue state, expected/persisted/failed counts, and timestamps. Persistence helpers catch database errors so observability writes do not replace the learner-facing result.",
    "Focused tests cover OTLP header parsing, disabled no-op behavior, opaque assessment IDs, scalar-only span attributes, and bounded-cardinality duration metrics. Production assessment-flow tests separately exercise lifecycle and persistence behavior.",
    "Current boundary: exporters run in batches and normal export failures occur outside the response path, but configure_telemetry itself has no protective try/except. A malformed initialization-time configuration can still prevent startup and should be guarded in a future hardening pass.",
]:
    insert_paragraph_before(doc, security_heading, text, "List Bullet")

# Expand the configuration/runbook/source references to match the actual implementation.
set_cell(
    config_table.cell(4, 1),
    "OTEL_ENABLED, OTEL_SERVICE_NAME, OTEL_EXPORTER_OTLP_ENDPOINT, OTEL_EXPORTER_OTLP_METRICS_ENDPOINT, OTEL_EXPORTER_OTLP_HEADERS, OTEL_METRIC_EXPORT_INTERVAL_MS, OTEL_SAMPLE_RATIO, OTEL_CAPTURE_CONTENT",
)
set_cell(config_table.cell(4, 2), "Optional OpenTelemetry/OpenInference traces and stage-duration metrics over OTLP HTTP")
set_para(
    find_para(doc, "OpenTelemetry export failures should not prevent application startup; the telemetry initializer is designed to fail soft."),
    "OpenTelemetry is disabled by default, and normal batch-export failures remain outside the learner-facing response path. Initialization is not currently guarded, so validate OTLP endpoints and headers before enabling it in production.",
)
set_para(
    find_para(doc, "Add explicit OpenTelemetry spans around Neo4j reads/writes and cache layers so dependency time is separated from route time."),
    "Add explicit OpenTelemetry spans around Neo4j reads/writes and cache layers so dependency time is separated from route time; also guard telemetry initialization so a malformed optional exporter cannot block application startup.",
)
set_labeled_cell(
    source_table.cell(0, 0),
    "Source files\n",
    "Primary implementation sources referenced by this document include backend/app/main.py, auth.py, config.py, database.py, observability.py, telemetry.py, assessment_observability.py, routers/curriculum.py, routers/test.py, routers/tutor.py, services/llm.py, services/generation_cache.py, services/assessment_attempts.py, services/semantic_clustering.py, services/clustering_runner.py, teacher_analytics.py, backend/.env.example, backend/tests/test_telemetry.py, backend/tests/test_assessment_observability.py, backend/tests/test_production_assessment_flow.py, and backend/tests/test_assessment_context.py. Unrelated evaluation, export, and promotion scripts are not used as behavioral evidence in this revision.",
)

# Preserve the accessibility metadata used by the prior revision.
for table in doc.tables:
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        tr_pr.append(header)

doc.save(DRAFT)
print(DRAFT)
